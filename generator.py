import os
import streamlit as st
from google import genai
from google.genai import types
import pypdf
from pptx import Presentation
import io
import zipfile

def get_api_key() -> str:
    """Retrieves API key from environment variables or Streamlit secrets."""
    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key:
        try:
            api_key = st.secrets["GEMINI_API_KEY"]
        except Exception:
            pass
    if not api_key:
        raise ValueError("GEMINI_API_KEY is missing from environment variables and Streamlit secrets.")
    return api_key

def extract_text_from_file(uploaded_file) -> str:
    """Extracts text content based on file extension."""
    filename = uploaded_file.name.lower()
    content = ""
    
    if filename.endswith(".pdf"):
        reader = pypdf.PdfReader(uploaded_file)
        for page in reader.pages:
            text = page.extract_text()
            if text:
                content += text + "\n"
    elif filename.endswith(".pptx"):
        prs = Presentation(uploaded_file)
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        content += paragraph.text + "\n"
    elif filename.endswith((".txt", ".md")):
        content = uploaded_file.getvalue().decode("utf-8")
    else:
        content = uploaded_file.getvalue().decode("utf-8", errors="ignore")
        
    return content

def generate_single_assessment(
    source_text: str,
    easy_count: int,
    med_count: int,
    hard_count: int,
    q_type: str,
    num_choices: int,
    include_answer_key: bool,
    output_format: str
) -> str:
    """Invokes Gemini API for a single source text with granular difficulty tiers and tagging."""
    api_key = get_api_key()
    client = genai.Client(api_key=api_key)
    
    total_questions = easy_count + med_count + hard_count
    
    if output_format == "Markdown (.md) with Tags":
        format_instruction = (
            "Output in clean Markdown (.md). Each question must be explicitly prefixed with "
            "a metadata tag indicating its tier (e.g., `[Difficulty: Easy]`, `[Difficulty: Medium]`, `[Difficulty: Hard]`) "
            "so automated Socratic parser scripts can easily ingest and index them."
        )
    elif output_format == "Structured XML (.xml)":
        format_instruction = (
            "Output in structured XML format using clean tags such as <assessment_bank>, <question_item>, "
            "<difficulty>, <question_text>, <options>, and <answer_key>."
        )
    else:
        format_instruction = "Output in plain text."

    choice_instruction = f"Each multiple choice question must have exactly {num_choices} distinct options (e.g., A through {chr(64 + num_choices)})." if q_type == "Multiple Choice" else ""
    answer_instruction = "Include a comprehensive answer key at the end of the document." if include_answer_key else "Do not include an answer key."
    
    prompt = f"""
    You are an expert academic assessment designer and university professor. 
    Analyze the following source document thoroughly. Automatically detect and match the academic level, technical rigor, and domain expertise present in the source text.

    SOURCE MATERIAL:
    """ + source_text[:30000] + f"""

    TASK REQUIREMENTS:
    - Generate a total of {total_questions} questions distributed across exact difficulty tiers:
      * Easy questions: {easy_count}
      * Medium questions: {med_count}
      * Hard questions: {hard_count}
    - Question Type: {q_type}
    - {choice_instruction}
    - {answer_instruction}
    - {format_instruction}
    - Ensure questions evaluate deep conceptual understanding and critical thinking directly grounded in the text.
    """

    response = client.models.generate_content(
        model="gemini-2.5-flash",
        contents=prompt,
        config=types.GenerateContentConfig(
            temperature=0.3,
            system_instruction="You produce rigorous, university-level academic assessments with strict structural tagging."
        )
    )
    
    return response.text

def generate_batch_assessments(
    uploaded_files,
    easy_count: int,
    med_count: int,
    hard_count: int,
    q_type: str,
    num_choices: int,
    include_answer_key: bool,
    output_format: str,
    progress_callback=None
) -> bytes:
    """Processes a queue of uploaded files and returns an in-memory ZIP archive of results."""
    zip_buffer = io.BytesIO()
    
    if output_format == "Structured XML (.xml)":
        file_extension = "xml"
    elif output_format == "Plain Text":
        file_extension = "txt"
    else:
        file_extension = "md"
    
    total_files = len(uploaded_files)
    
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
        for idx, uploaded_file in enumerate(uploaded_files):
            if progress_callback:
                progress_callback(idx, total_files, uploaded_file.name)
                
            source_text = extract_text_from_file(uploaded_file)
            assessment_text = generate_single_assessment(
                source_text=source_text,
                easy_count=easy_count,
                med_count=med_count,
                hard_count=hard_count,
                q_type=q_type,
                num_choices=num_choices,
                include_answer_key=include_answer_key,
                output_format=output_format
            )
            
            base_name = os.path.splitext(uploaded_file.name)[0]
            out_filename = f"{base_name}_question_bank.{file_extension}"
            zip_file.writestr(out_filename, assessment_text)
            
    zip_buffer.seek(0)
    return zip_buffer.getvalue()
